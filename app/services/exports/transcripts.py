"""Interview exports: the message table (CSV, Excel) and readable transcripts
(Word, PDF), each as one file or as a ZIP with one file per interview.

The transcript formats render from the same `TranscriptView`, so a Word file
and a PDF of one interview say the same thing; only the layout differs. PDF
goes through the shared HTML templates (see `rendering`). Word has no such
route -- nothing turns styled HTML into a decent .docx -- so it is built
directly with python-docx.
"""

import datetime
import io
import re
import zipfile
from collections.abc import Callable, Sequence
from dataclasses import dataclass, field
from typing import Any, Literal
from zoneinfo import ZoneInfo

import polars as pl
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.text.paragraph import Paragraph
from markupsafe import Markup
from pydantic import UUID4

from ainterviewer.types import InterviewStatus, MessageRole

from ...db.utils import fix_nested_columns, messages_to_dataframe, write_messages_xlsx
from ...schemas.interviews import InterviewSummaryPublic, MessagePublic
from ...types import ExternalParam, InterviewType
from ...utils import as_aware
from . import markup
from .markup import Span
from .rendering import render_pdf

ExportFormat = Literal["csv", "xlsx", "docx", "pdf"]

MEDIA_TYPES: dict[str, str] = {
    "csv": "text/csv",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pdf": "application/pdf",
    "zip": "application/zip",
}

STATUS_LABELS = {
    InterviewStatus.COMPLETED: "Complete",
    InterviewStatus.ACTIVE: "Active",
    InterviewStatus.INACTIVE: "Inactive",
}

TYPE_LABELS = {
    InterviewType.MANUAL_TEST: "Manual test",
    InterviewType.SYNTHETIC_TEST: "Synthetic test",
    InterviewType.DISTRIBUTED: "Distributed",
}

# The platform's names for survey item types (`frontend/src/lib/survey/itemTypes.ts`).
SURVEY_LABELS = {
    "radio": "Single choice",
    "checkbox": "Multiple choice",
    "likert": "Likert",
    "slider": "Slider",
    "number": "Number",
    "date": "Date",
    "datetime": "Date & time",
    "time": "Time",
}

# Characters XML 1.0 forbids. python-docx raises on them, and they mean
# nothing in a transcript anyway.
_XML_INVALID = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
_UNSAFE_FILENAME = re.compile(r"[^A-Za-z0-9._-]+")


@dataclass(frozen=True)
class InterviewExport:
    """One interview, with everything any export format needs."""

    summary: InterviewSummaryPublic
    messages: Sequence[MessagePublic]
    external_params: dict[str, Any] = field(default_factory=dict)


@dataclass(frozen=True)
class SurveyAnswer:
    """A survey answer with the options it was chosen from, as the platform
    shows it: an "Agree" means little without the scale around it."""

    kind: str
    # (option, chosen); empty for items that are not a choice.
    options: list[tuple[str, bool]]
    answer: str
    # A slider's range, `0 (Not at all) – 10 (Very)`.
    scale: str | None = None


@dataclass(frozen=True)
class Turn:
    """One row of the transcript: a message, a control token (`Probe
    complete`) shown as a pill, or the start of a guide section."""

    kind: Literal["message", "token", "section"]
    # The message as plain text, or the token's or section's label.
    text: str
    spans: list[Span] = field(default_factory=list)
    speaker: str = ""
    time: str = ""
    is_participant: bool = False
    # The guide number, `3.2.1`, on an interviewer's question.
    label: str | None = None
    survey: SurveyAnswer | None = None

    @property
    def html(self) -> Markup:
        return markup.to_html(self.spans)


@dataclass(frozen=True)
class TranscriptView:
    title: str
    meta: list[tuple[str, str]]
    turns: list[Turn]


@dataclass(frozen=True)
class ExportFile:
    content: bytes
    filename: str
    media_type: str


def _clean(text: str) -> str:
    return _XML_INVALID.sub("", text)


def _format_time(value: datetime.datetime | None, tz: ZoneInfo, fmt: str) -> str:
    if value is None:
        return ""
    return as_aware(value).astimezone(tz).strftime(fmt)


def _image_note(message: MessagePublic) -> str | None:
    images = message.image if isinstance(message.image, list) else [message.image]
    labels = [
        getattr(image, "alt", None) or getattr(image, "description", None) or ""
        for image in images
        if image is not None
    ]
    if not labels:
        return None
    return "\n".join(f"[Image{': ' + label if label else ''}]" for label in labels)


def _question_label(message: MessagePublic) -> str | None:
    # A section's intro or outro has no question number worth showing.
    if message.section is None or message.main_question is None:
        return None
    label = f"{message.section + 1}.{message.main_question + 1}"
    # sub_question 0 is the main question itself, above that a probe.
    if message.sub_question:
        label += f".{message.sub_question}"
    return label


def _checkbox_choices(answer: str, options: list[str]) -> tuple[set[str], list[str]]:
    """Split a checkbox answer back into its choices.

    The answer is the choices joined by ", ", and an option may contain a
    comma itself, so the options are matched first, longest first, and only
    the rest is split: that is what the participant wrote under "Other".
    """
    rest = answer
    chosen: set[str] = set()
    for option in sorted(options, key=len, reverse=True):
        pattern = rf"(?:^|,)\s*{re.escape(option)}\s*(?=,|$)"
        if re.search(pattern, rest):
            chosen.add(option)
            rest = re.sub(pattern, "", rest, count=1)
    return chosen, [part.strip() for part in rest.split(",") if part.strip()]


def _survey_answer(item: Any, answer: str) -> SurveyAnswer:
    """Mark the chosen options, the way the platform's read-only item does.

    A checkbox answer is its choices joined by commas; an answer that is not
    one of the options is what the participant wrote under "Other".
    """
    kind = SURVEY_LABELS.get(item.type, str(item.type))
    options: list[str] = list(getattr(item, "options", None) or [])
    if not options:
        scale = None
        if item.type == "slider":
            ends = [
                f"{bound} ({label})" if label else f"{bound}"
                for bound, label in (
                    (item.min, item.min_label),
                    (item.max, item.max_label),
                )
            ]
            scale = " – ".join(ends)
        return SurveyAnswer(kind=kind, options=[], answer=answer, scale=scale)

    if item.type == "checkbox":
        chosen, others = _checkbox_choices(answer, options)
    else:
        chosen = {answer.strip()} & set(options)
        others = [] if chosen or not answer.strip() else [answer.strip()]
    marked = [(option, option in chosen) for option in options]
    marked += [(f"Other: {other}", True) for other in others]
    return SurveyAnswer(kind=kind, options=marked, answer=answer)


def build_view(interview: InterviewExport, tz: ZoneInfo) -> TranscriptView:
    """The transcript as the platform shows it: who said what, and when,
    grouped by guide section.

    System messages and questions skipped by a condition are left out -- the
    participant never saw either.
    """
    summary = interview.summary
    title = summary.pid or summary.test_name or f"Interview {str(summary.id)[:8]}"

    meta: list[tuple[str, str]] = [("Interview ID", str(summary.id))]
    if summary.pid:
        meta.append(("PID", summary.pid))
    if summary.test_name:
        meta.append(("Test", summary.test_name))
    if summary.type != InterviewType.DISTRIBUTED:
        meta.append(("Type", TYPE_LABELS[summary.type]))
    meta += [
        ("Language", summary.language),
        ("Status", STATUS_LABELS.get(summary.status, str(summary.status))),
        ("Started", _format_time(summary.created_at, tz, "%d/%m/%Y %H:%M")),
    ]
    if summary.last_updated:
        meta.append(
            ("Last updated", _format_time(summary.last_updated, tz, "%d/%m/%Y %H:%M"))
        )
    meta.append(("Time zone", str(tz)))
    for name, value in interview.external_params.items():
        # A `pid` param repeats the PID row above.
        if value in (None, "") or (name == "pid" and value == summary.pid):
            continue
        meta.append((name, str(value)))

    turns: list[Turn] = []
    section: int | None = None
    pending_survey: Any = None
    for message in sorted(interview.messages, key=lambda m: m.message_id):
        if message.role == MessageRole.SYSTEM or message.skipped_by_condition:
            continue

        if message.section is not None and message.section != section:
            section = message.section
            turns.append(Turn(kind="section", text=f"Section {section + 1}"))

        is_participant = message.role == MessageRole.USER
        content = _clean(message.content.strip())
        # Guide markup is the interviewer's; a participant's `<b>` is text.
        spans = markup.parse(content, markup=not is_participant)

        if token := markup.only_token(spans):
            turns.append(Turn(kind="token", text=token, is_participant=is_participant))
            if is_participant:
                pending_survey = None
            continue

        if note := _image_note(message):
            spans = [*spans, Span(("\n" if spans else "") + _clean(note))]

        survey = None
        if is_participant and pending_survey is not None:
            survey = _survey_answer(pending_survey, content)
            pending_survey = None
        elif not is_participant:
            pending_survey = message.survey_item

        if not spans and survey is None:
            continue
        turns.append(
            Turn(
                kind="message",
                text=markup.plain(spans),
                spans=spans,
                speaker="Participant" if is_participant else "Interviewer",
                time=_format_time(message.created_at, tz, "%H:%M"),
                is_participant=is_participant,
                label=None if is_participant else _question_label(message),
                survey=survey,
            )
        )

    return TranscriptView(title=_clean(title), meta=meta, turns=turns)


# ------------------------------------------------------------------ formats


def _table(
    interviews: Sequence[InterviewExport], external_params: Sequence[ExternalParam]
) -> pl.DataFrame:
    return messages_to_dataframe(
        [message for interview in interviews for message in interview.messages],
        external_params=external_params,
        external_param_values={
            interview.summary.id: interview.external_params for interview in interviews
        },
    )


def render_csv(
    interviews: Sequence[InterviewExport],
    external_params: Sequence[ExternalParam],
    tz: ZoneInfo,
) -> bytes:
    with io.BytesIO() as stream:
        fix_nested_columns(_table(interviews, external_params)).write_csv(stream)
        return stream.getvalue()


def render_xlsx(
    interviews: Sequence[InterviewExport],
    external_params: Sequence[ExternalParam],
    tz: ZoneInfo,
) -> bytes:
    with io.BytesIO() as stream:
        write_messages_xlsx(_table(interviews, external_params), stream)
        return stream.getvalue()


def render_transcripts_pdf(
    interviews: Sequence[InterviewExport],
    external_params: Sequence[ExternalParam],
    tz: ZoneInfo,
) -> bytes:
    return render_pdf(
        "transcripts.html.j2",
        {"transcripts": [build_view(interview, tz) for interview in interviews]},
    )


_MUTED = RGBColor(0x6B, 0x72, 0x80)
_ACCENT = RGBColor(0x19, 0x68, 0x58)
_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
_INTERVIEWER_FILL = "EEEEEE"
_PARTICIPANT_FILL = "196858"
_SURVEY_FILL = "E5E7EB"


def _shade(paragraph: Paragraph, fill: str) -> None:
    """Fill the paragraph's box, which is as close to a bubble as Word gets."""
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    paragraph._p.get_or_add_pPr().append(shading)


def _add_spans(paragraph: Paragraph, spans: Sequence[Span], color=None) -> None:
    for span in spans:
        run = paragraph.add_run(span.text)
        run.bold = span.token or "bold" in span.marks or "link" in span.marks or None
        run.italic = "italic" in span.marks or None
        run.underline = "underline" in span.marks or "link" in span.marks or None
        run.font.strike = "strike" in span.marks or None
        if "code" in span.marks:
            run.font.name = "Consolas"
        if color is not None:
            run.font.color.rgb = color


def _bubble(document, turn: Turn) -> list[Paragraph]:
    """The turn's paragraphs: shaded, and pushed to its side of the page.

    The first line is the platform's header row: the guide number (or the
    survey item's type) on the left, the time against the right edge.
    """
    participant = turn.is_participant
    paragraphs: list[Paragraph] = []
    section = document.sections[-1]
    text_width = section.page_width - section.left_margin - section.right_margin

    def add() -> Paragraph:
        paragraph = document.add_paragraph()
        fmt = paragraph.paragraph_format
        fmt.space_before = fmt.space_after = Pt(0)
        fmt.keep_with_next = True
        fmt.keep_together = True
        if participant:
            fmt.left_indent = Pt(96)
        else:
            fmt.right_indent = Pt(96)
        paragraphs.append(paragraph)
        return paragraph

    def header(title: str | None, color: RGBColor) -> None:
        paragraph = add()
        # Tab stops count from the margin, not the indent, so the bubble's
        # right edge is the text width less whatever indent is on that side.
        right_edge = text_width - (paragraph.paragraph_format.right_indent or 0)
        paragraph.paragraph_format.tab_stops.add_tab_stop(
            right_edge, WD_TAB_ALIGNMENT.RIGHT
        )
        for text, bold in ((title or "", True), (f"\t{turn.time}", None)):
            run = paragraph.add_run(text)
            run.bold = bold
            run.font.size = Pt(7.5)
            run.font.color.rgb = color

    if turn.survey is not None:
        header(turn.survey.kind, _MUTED)
        if turn.survey.options:
            for option, chosen in turn.survey.options:
                run = add().add_run(f"{'☑' if chosen else '☐'}  {option}")
                run.bold = chosen or None
                if not chosen:
                    run.font.color.rgb = _MUTED
        else:
            add().add_run(turn.survey.answer).bold = True
            if turn.survey.scale:
                scale = add().add_run(turn.survey.scale)
                scale.font.size = Pt(7.5)
                scale.font.color.rgb = _MUTED
        fill = _SURVEY_FILL
    else:
        header(turn.label, _WHITE if participant else _MUTED)
        # A run's text setter turns "\n" into a line break.
        _add_spans(add(), turn.spans, _WHITE if participant else None)
        fill = _PARTICIPANT_FILL if participant else _INTERVIEWER_FILL

    for paragraph in paragraphs:
        _shade(paragraph, fill)
    paragraphs[-1].paragraph_format.space_after = Pt(8)
    return paragraphs


def render_transcripts_docx(
    interviews: Sequence[InterviewExport],
    external_params: Sequence[ExternalParam],
    tz: ZoneInfo,
) -> bytes:
    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)

    for index, interview in enumerate(interviews):
        view = build_view(interview, tz)
        if index:
            document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

        eyebrow = document.add_paragraph().add_run("INTERVIEW TRANSCRIPT")
        eyebrow.bold = True
        eyebrow.font.size = Pt(7.5)
        eyebrow.font.color.rgb = _MUTED
        heading = document.add_heading(view.title, level=1)
        for run in heading.runs:
            run.font.color.rgb = _ACCENT

        table = document.add_table(rows=0, cols=2)
        for label, value in view.meta:
            cells = table.add_row().cells
            for cell in cells:
                cell.paragraphs[0].paragraph_format.space_after = Pt(0)
            label_run = cells[0].paragraphs[0].add_run(_clean(label))
            label_run.font.color.rgb = _MUTED
            label_run.font.size = Pt(8.5)
            cells[1].paragraphs[0].add_run(_clean(value)).font.size = Pt(8.5)
        document.add_paragraph()

        if not view.turns:
            document.add_paragraph().add_run("No messages.").italic = True

        for turn in view.turns:
            if turn.kind == "section":
                paragraph = document.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.space_before = Pt(12)
                paragraph.paragraph_format.keep_with_next = True
                run = paragraph.add_run(turn.text.upper())
                run.bold = True
                run.font.size = Pt(7.5)
                run.font.color.rgb = _MUTED
                continue
            if turn.kind == "token":
                paragraph = document.add_paragraph()
                paragraph.alignment = (
                    WD_ALIGN_PARAGRAPH.RIGHT
                    if turn.is_participant
                    else WD_ALIGN_PARAGRAPH.CENTER
                )
                run = paragraph.add_run(turn.text)
                run.bold = True
                run.font.size = Pt(8)
                run.font.color.rgb = _MUTED
                continue

            _bubble(document, turn)

    with io.BytesIO() as stream:
        document.save(stream)
        return stream.getvalue()


Renderer = Callable[
    [Sequence[InterviewExport], Sequence[ExternalParam], ZoneInfo], bytes
]

RENDERERS: dict[str, Renderer] = {
    "csv": render_csv,
    "xlsx": render_xlsx,
    "docx": render_transcripts_docx,
    "pdf": render_transcripts_pdf,
}


def _file_stem(summary: InterviewSummaryPublic) -> str:
    stem = summary.pid or summary.test_name or str(summary.id)
    return _UNSAFE_FILENAME.sub("_", stem).strip("._") or str(summary.id)


def export_interviews(
    interviews: Sequence[InterviewExport],
    *,
    project_id: UUID4,
    format: ExportFormat,
    archive: bool,
    external_params: Sequence[ExternalParam],
    tz: ZoneInfo,
) -> ExportFile:
    """One file holding every interview, or a ZIP with one file per interview.

    Archive entries are named after the participant (or test, or interview
    id); two interviews of the same participant get the id appended rather
    than overwriting each other.
    """
    render = RENDERERS[format]

    if not archive:
        if len(interviews) == 1:
            stem = _file_stem(interviews[0].summary)
        else:
            kind = "messages" if format in ("csv", "xlsx") else "transcripts"
            stem = f"{project_id}-{kind}"
        return ExportFile(
            content=render(interviews, external_params, tz),
            filename=f"{stem}.{format}",
            media_type=MEDIA_TYPES[format],
        )

    used: set[str] = set()
    with io.BytesIO() as stream:
        with zipfile.ZipFile(stream, "w", zipfile.ZIP_DEFLATED) as zf:
            for interview in interviews:
                stem = _file_stem(interview.summary)
                if stem in used:
                    stem = f"{stem}-{str(interview.summary.id)[:8]}"
                used.add(stem)
                zf.writestr(
                    f"{stem}.{format}", render([interview], external_params, tz)
                )
        content = stream.getvalue()

    return ExportFile(
        content=content,
        filename=f"{project_id}-{format}.zip",
        media_type=MEDIA_TYPES["zip"],
    )
