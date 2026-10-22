"""Interview exports in every format, as one file and as a ZIP."""

import datetime
import io
import uuid
import zipfile
from zoneinfo import ZoneInfo

import pytest
from docx import Document
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

from ainterviewer.interview_guides import InterviewGuide
from ainterviewer.interview_guides.survey_items import CheckboxItem
from ainterviewer.types import InterviewStatus, MessageRole
from app.db.crud import InterviewDataBase
from app.db.tables import (
    Base,
    InterviewTable,
    ProjectFolderTable,
    ProjectTable,
    UserTable,
)
from app.schemas.interviews import InterviewSummaryPublic, MessagePublic
from app.services.exports import markup
from app.services.exports.transcripts import (
    InterviewExport,
    _survey_answer,
    build_view,
    export_interviews,
)
from app.types import InterviewType

PROJECT_ID = uuid.uuid4()
UTC = ZoneInfo("UTC")
STARTED = datetime.datetime(2026, 9, 8, 10, 48, tzinfo=datetime.UTC)


def _message(interview_id, message_id, role, content, **extra) -> MessagePublic:
    return MessagePublic(
        id=uuid.uuid4(),
        message_id=message_id,
        content=content,
        role=role,
        interview_id=interview_id,
        project_id=PROJECT_ID,
        created_at=STARTED + datetime.timedelta(minutes=message_id),
        interview_type=InterviewType.DISTRIBUTED,
        **extra,
    )


def _interview(pid: str | None) -> InterviewExport:
    interview_id = uuid.uuid4()
    return InterviewExport(
        summary=InterviewSummaryPublic(
            id=interview_id,
            language="DA",
            status=InterviewStatus.COMPLETED,
            type=InterviewType.DISTRIBUTED,
            created_at=STARTED,
            n_messages=4,
            pid=pid,
        ),
        # Out of order on purpose: transcripts read in message order.
        messages=[
            _message(interview_id, 3, MessageRole.USER, "Selvfølgelig, æøå 🙂"),
            _message(interview_id, 0, MessageRole.SYSTEM, "internal prompt"),
            _message(
                interview_id, 1, MessageRole.ASSISTANT, "Hej! Tak fordi du\nvil deltage"
            ),
            _message(
                interview_id,
                2,
                MessageRole.ASSISTANT,
                "Skipped",
                skipped_by_condition=True,
            ),
        ],
        external_params={"cohort": "B"},
    )


def _export(interviews, format, archive):
    return export_interviews(
        interviews,
        project_id=PROJECT_ID,
        format=format,
        archive=archive,
        external_params=[],
        tz=UTC,
    )


def test_view_reads_like_the_conversation():
    view = build_view(_interview("P1"), UTC)
    assert view.title == "P1"
    assert [(turn.speaker, turn.text) for turn in view.turns] == [
        ("Interviewer", "Hej! Tak fordi du\nvil deltage"),
        ("Participant", "Selvfølgelig, æøå 🙂"),
    ]
    assert view.turns[0].time == "10:49"
    assert ("cohort", "B") in view.meta


def test_view_times_follow_the_requested_zone():
    view = build_view(_interview("P1"), ZoneInfo("Europe/Copenhagen"))
    assert view.turns[0].time == "12:49"


@pytest.mark.parametrize("format", ["csv", "xlsx", "docx", "pdf"])
def test_single_file_holds_every_interview(format):
    export = _export([_interview("P1"), _interview("P2")], format, archive=False)
    assert export.filename.endswith(f".{format}")
    assert export.content
    if format == "pdf":
        assert export.content.startswith(b"%PDF")
    if format == "csv":
        assert "æøå".encode() in export.content


def test_single_interview_is_named_after_its_participant():
    export = _export([_interview("P1")], "pdf", archive=False)
    assert export.filename == "P1.pdf"


def test_docx_carries_the_transcript_text():
    export = _export([_interview("P1")], "docx", archive=False)
    text = "\n".join(p.text for p in Document(io.BytesIO(export.content)).paragraphs)
    assert "Selvfølgelig, æøå 🙂" in text
    assert "internal prompt" not in text
    assert "Skipped" not in text


@pytest.mark.parametrize("format", ["csv", "xlsx", "docx", "pdf"])
def test_archive_has_one_file_per_interview(format):
    same_pid = [_interview("P1"), _interview("P1"), _interview(None)]
    export = _export(same_pid, format, archive=True)
    assert export.media_type == "application/zip"
    names = zipfile.ZipFile(io.BytesIO(export.content)).namelist()
    assert len(names) == 3
    # The second P1 must not overwrite the first.
    assert len(set(names)) == 3
    assert names[0] == f"P1.{format}"
    assert all(name.endswith(f".{format}") for name in names)


def test_summaries_stay_inside_the_project():
    engine = create_engine("sqlite://")
    Base.metadata.create_all(engine)
    with Session(engine) as session:
        user = UserTable(email="a@example.com", password="x", first_name="A")
        folder = ProjectFolderTable(title="f")
        project = ProjectTable(title="p", folder=folder, owner=user)
        other = ProjectTable(title="q", folder=folder, owner=user)
        mine, theirs = (
            InterviewTable(
                project=owner,
                interview_guide=InterviewGuide(),
                type=InterviewType.DISTRIBUTED,
            )
            for owner in (project, other)
        )
        session.add_all([user, folder, project, other, mine, theirs])
        session.commit()

        summaries = InterviewDataBase(session).interviews.get_interview_summaries(
            project.id, [mine.id, theirs.id]
        )
        assert list(summaries) == [mine.id]


@pytest.mark.parametrize("format", ["csv", "xlsx", "docx", "pdf"])
@pytest.mark.parametrize("archive", [False, True])
def test_an_interview_without_messages_still_exports(format, archive):
    empty = _interview("P0")
    empty = InterviewExport(summary=empty.summary, messages=[])
    assert _export([empty], format, archive=archive).content


def test_guide_markup_follows_the_platform_allowlist():
    spans = markup.parse(
        'A <b>bold</b> <a href="https://x.org">link</a>, <a href="javascript:x()">no</a>'
        "<script>alert(1)</script> and <blink>",
        markup=True,
    )
    html = str(markup.to_html(spans))
    assert "<b>bold</b>" in html
    # Styled as a link, but not one: no target ends up in the document.
    assert '<span class="link">link</span>' in html
    assert "<a" not in html and "x.org" not in html
    assert "javascript" not in html and ">no<" not in html.replace("<b>", "")
    assert "alert" not in html
    # Outside the allowlist stays visible, as on the platform.
    assert "&lt;blink&gt;" in html


def test_participant_text_is_never_markup():
    spans = markup.parse("I <b>typed</b> this", markup=False)
    assert markup.plain(spans) == "I <b>typed</b> this"
    assert "&lt;b&gt;" in str(markup.to_html(spans))


def test_tokens_become_labels_even_inside_markup():
    spans = markup.parse("Thanks <i>a lot</i>! <|endofinterview|>", markup=True)
    assert markup.plain(spans) == "Thanks a lot! [Interview complete]"
    assert markup.only_token(markup.parse(" <|skipquestion|> ", markup=False)) == (
        "Question skipped"
    )


def test_view_divides_sections_and_numbers_questions():
    interview_id = uuid.uuid4()
    messages = [
        _message(interview_id, 0, MessageRole.ASSISTANT, "Intro", section=0),
        _message(
            interview_id,
            1,
            MessageRole.ASSISTANT,
            "Tools?",
            section=0,
            main_question=1,
            sub_question=0,
            survey_item=CheckboxItem(options=["Email", "Chat, voice", "Paper"]),
        ),
        _message(interview_id, 2, MessageRole.USER, "Email, Whiteboard", section=0),
        _message(interview_id, 3, MessageRole.ASSISTANT, "<|endofsection|>", section=0),
        _message(
            interview_id,
            4,
            MessageRole.ASSISTANT,
            "Why?",
            section=1,
            main_question=0,
            sub_question=2,
        ),
        _message(interview_id, 5, MessageRole.USER, "<|skipquestion|>", section=1),
    ]
    interview = _interview("P1")
    view = build_view(
        InterviewExport(
            summary=interview.summary,
            messages=messages,
            external_params={"pid": "P1", "cohort": "B"},
        ),
        UTC,
    )
    assert [(turn.kind, turn.text) for turn in view.turns] == [
        ("section", "Section 1"),
        ("message", "Intro"),
        ("message", "Tools?"),
        ("message", "Email, Whiteboard"),
        ("token", "Section complete"),
        ("section", "Section 2"),
        ("message", "Why?"),
        ("token", "Question skipped"),
    ]
    intro, question, answer = view.turns[1:4]
    assert intro.label is None
    assert question.label == "1.2"
    assert view.turns[6].label == "2.1.2"
    assert answer.survey is not None
    assert answer.survey.kind == "Multiple choice"
    assert answer.survey.options == [
        ("Email", True),
        ("Chat, voice", False),
        ("Paper", False),
        ("Other: Whiteboard", True),
    ]
    assert view.turns[7].is_participant
    # The pid param would only repeat the PID row.
    assert [label for label, _ in view.meta].count("pid") == 0
    assert _survey_answer(
        CheckboxItem(options=["Email", "Chat, voice"]), "Chat, voice, Email"
    ).options == [("Email", True), ("Chat, voice", True)]
    assert ("cohort", "B") in view.meta
